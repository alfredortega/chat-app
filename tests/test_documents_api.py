"""
Backend tests for the unified Markdown document API
(``/api/conversations/<id>/documents`` list/read/update).

Covers output-directory + conversation-upload backends, the output-dir jail
(traversal, symlinks, absolute paths, null bytes), optimistic concurrency,
atomic writes, ConvFile metadata updates + cache invalidation, managed project
artifacts going through the human-authored review workflow, ``document_created``
SSE events after model writes, and the recent-edit context note.
"""

import json
import os
import tempfile
import shutil

import pytest

import database as db_module
import documents as documents_module
import file_handler
from app import create_app

# Directories created during a test, removed after it (autouse fixture below).
_AUTO_CLEANUP: list[str] = []


@pytest.fixture(autouse=True)
def _cleanup_temp_dirs():
    yield
    for path in _AUTO_CLEANUP:
        shutil.rmtree(path, ignore_errors=True)
    _AUTO_CLEANUP.clear()


def _app(tmp_db):
    return create_app(config={
        "SQLALCHEMY_DATABASE_URI": tmp_db,
        "ENCRYPTION_KEY": "test-encryption-key-must-be-32-bytes-long-!!!!",
        "SKIP_DOTENV_WRITE": "1",
    })


class _Context:
    """Holds an app + a ready conversation + a temp output/session dir."""

    def __init__(self, app, output_dir=None, local_access=True):
        self.app = app
        self.tmp_dir = output_dir or tempfile.mkdtemp()
        with app.app_context():
            db_module.set_setting("allow_local_file_access", "1" if local_access else "0")
            db_module.set_setting("output_dir", self.tmp_dir)
            endpoint = db_module.create_endpoint(
                "fake", "http://fase.openai.test/v1", api_key="f",
                default_model="test-model", is_default=True,
            )
            conv = db_module.create_conversation("Doc test", "test-model", endpoint_id=endpoint["id"])
            self.conv_id = conv["id"]

    def close(self):
        shutil.rmtree(self.tmp_dir, ignore_errors=True)


def _make(app, output_dir=None, local_access=True, request=None):
    ctx = _Context(app, output_dir=output_dir, local_access=local_access)
    _AUTO_CLEANUP.append(ctx.tmp_dir)
    return ctx


def _write_output(ctx, rel, content):
    path = os.path.join(ctx.tmp_dir, *rel.split("/"))
    os.makedirs(os.path.dirname(path), exist_ok=True)
    with open(path, "w", encoding="utf-8") as fh:
        fh.write(content)
    return path


def _add_upload(ctx, name, content, request=None):
    root = tempfile.mkdtemp()
    _AUTO_CLEANUP.append(root)
    disk = os.path.join(root, name)
    with open(disk, "w", encoding="utf-8") as fh:
        fh.write(content)
    with ctx.app.app_context():
        return db_module.add_conv_file(
            conversation_id=ctx.conv_id,
            original_name=name,
            disk_path=disk,
            size_bytes=os.path.getsize(disk),
            char_count=len(content),
        )


def _hash_sha256(content: str) -> str:
    import hashlib
    return "sha256:" + hashlib.sha256(content.encode("utf-8")).hexdigest()


@pytest.fixture
def ctx(tmp_db, request):
    """A conversation with the output dir + uploads configured (local access on)."""
    app = _app(tmp_db)
    return _make(app, request=request)


class TestListDocuments:
    def test_lists_only_markdown_in_output_dir(self, tmp_db):
        app = _app(tmp_db)
        c = _make(app, local_access=True)
        _write_output(c, "report.md", "# Report\n\nBody")
        _write_output(c, "notes.txt", "not markdown")
        _write_output(c, "sub/readme.markdown", "# Readme")
        _write_output(c, ".hidden.md", "# Hidden")
        with app.test_client() as client:
            resp = client.get(f"/api/conversations/{c.conv_id}/documents")
        assert resp.status_code == 200
        docs = resp.get_json()["documents"]
        names = sorted(d["name"] for d in docs)
        assert names == ["readme.markdown", "report.md"]
        rep = next(d for d in docs if d["name"] == "report.md")
        assert rep["kind"] == "output"
        assert rep["id"] == "output:report.md"
        assert rep["editable"] is True
        assert rep["managed_artifact"] is False
        assert rep["size_bytes"] == len("# Report\n\nBody")
        assert rep["modified_at"]
        assert rep["content_hash"].startswith("sha256:")

    def test_includes_markdown_uploads_but_not_other_uploads(self, ctx):
        _add_upload(ctx, "uploaded.md", "# From upload")
        _add_upload(ctx, "notes.txt", "ignore me")
        app = ctx.app
        with app.test_client() as client:
            resp = client.get(f"/api/conversations/{ctx.conv_id}/documents")
        docs = resp.get_json()["documents"]
        uploads = [d for d in docs if d["kind"] == "upload"]
        assert [d["name"] for d in uploads] == ["uploaded.md"]
        assert uploads[0]["id"].startswith("upload:")
        assert uploads[0]["file_id"] is not None

    def test_dedupes_output_and_upload_with_same_name(self, ctx):
        _write_output(ctx, "report.md", "# output version")
        _add_upload(ctx, "report.md", "# upload version")
        app = ctx.app
        with app.test_client() as client:
            resp = client.get(f"/api/conversations/{ctx.conv_id}/documents")
        docs = resp.get_json()["documents"]
        same = [d for d in docs if d["name"] == "report.md"]
        assert len(same) == 1
        assert same[0]["kind"] == "output"  # output wins the name

    def test_404_when_conversation_missing(self, tmp_db):
        app = _app(tmp_db)
        with app.test_client() as client:
            resp = client.get("/api/conversations/9999/documents")
        assert resp.status_code == 404


class TestOutputSecurity:
    def make_report(self, ctx):
        _write_output(ctx, "report.md", "# Report")
        return f"/api/conversations/{ctx.conv_id}/documents"

    def test_traversal_rejected(self, ctx):
        base = self.make_report(ctx)
        app = ctx.app
        with app.test_client() as client:
            for bad in ("output:../secret.md", "output:sub/../secret.md", "output:/etc/passwd"):
                resp = client.get(f"{base}/{bad}")
                assert resp.status_code in (400, 404), bad
            resp = client.get(f"{base}/output:..%2F..%2Fsecret.md")
            assert resp.status_code in (400, 404)

    def test_null_byte_rejected(self, ctx):
        base = self.make_report(ctx)
        app = ctx.app
        with app.test_client() as client:
            resp = client.get(base + "/output:rep%00ort.md")
            assert resp.status_code in (400, 404)

    def test_symlink_outside_root_skipped_in_list_and_rejected_on_read(self, ctx):
        outside = tempfile.mkdtemp()
        _AUTO_CLEANUP.append(outside)
        with open(os.path.join(outside, "leak.md"), "w", encoding="utf-8") as fh:
            fh.write("# secret")
        os.symlink(os.path.join(outside, "leak.md"), os.path.join(ctx.tmp_dir, "leak.md"))
        app = ctx.app
        with app.test_client() as client:
            resp = client.get(f"/api/conversations/{ctx.conv_id}/documents")
        names = [d["name"] for d in resp.get_json()["documents"]]
        assert "leak.md" not in names
        with app.test_client() as client:
            resp = client.get(f"/api/conversations/{ctx.conv_id}/documents/output:leak.md")
        assert resp.status_code == 400

    def test_upload_ownership_enforced(self, ctx):
        app = ctx.app
        up = _add_upload(ctx, "mine.md", "# Mine")
        # Second conversation must not see this conversation's upload.
        with app.app_context():
            other = db_module.create_conversation("Other", "m")
        with app.test_client() as client:
            resp = client.get(f"/api/conversations/{other['id']}/documents/upload:{up['id']}")
        assert resp.status_code == 404

    def test_unknown_document_id_shape(self, ctx):
        app = ctx.app
        with app.test_client() as client:
            resp = client.get(f"/api/conversations/{ctx.conv_id}/documents/foobar:1")
        assert resp.status_code == 400


class TestRead:
    def test_read_output_document_raw(self, ctx):
        _write_output(ctx, "report.md", "# Report\n\nline one\n")
        app = ctx.app
        with app.test_client() as client:
            resp = client.get(f"/api/conversations/{ctx.conv_id}/documents/output:report.md")
        assert resp.status_code == 200
        data = resp.get_json()
        assert data["content"] == "# Report\n\nline one\n"
        assert data["kind"] == "output"
        assert data["content_hash"] == _hash_sha256("# Report\n\nline one\n")
        assert data["editable"] is True

    def test_read_upload_document_raw(self, ctx):
        up = _add_upload(ctx, "notes.md", "# Notes\n")
        app = ctx.app
        with app.test_client() as client:
            resp = client.get(f"/api/conversations/{ctx.conv_id}/documents/upload:{up['id']}")
        assert resp.status_code == 200
        assert resp.get_json()["content"] == "# Notes\n"

    def test_read_missing_document_404(self, ctx):
        app = ctx.app
        with app.test_client() as client:
            resp = client.get(f"/api/conversations/{ctx.conv_id}/documents/output:missing.md")
        assert resp.status_code == 404


class TestUpdateOutput:
    def test_successful_update(self, ctx):
        _write_output(ctx, "report.md", "# Old")
        app = ctx.app
        new = "# New content"
        with app.test_client() as client:
            resp = client.put(
                f"/api/conversations/{ctx.conv_id}/documents/output:report.md",
                json={"content": new, "expected_hash": _hash_sha256("# Old")},
            )
        assert resp.status_code == 200
        data = resp.get_json()
        assert data["success"] is True
        assert data["content_hash"] == _hash_sha256(new)
        assert data["size_bytes"] == len(new)
        with open(os.path.join(ctx.tmp_dir, "report.md"), encoding="utf-8") as fh:
            assert fh.read() == new

    def test_hash_mismatch_409_keeps_local_and_returns_server(self, ctx):
        _write_output(ctx, "report.md", "# Old")
        app = ctx.app
        with app.test_client() as client:
            resp = client.put(
                f"/api/conversations/{ctx.conv_id}/documents/output:report.md",
                json={"content": "# Local edit", "expected_hash": "sha256:wrong"},
            )
        assert resp.status_code == 409
        payload = resp.get_json()
        assert payload["error"] == "document_changed"
        assert payload["current_hash"] == _hash_sha256("# Old")
        assert payload["current_content"] == "# Old"
        with open(os.path.join(ctx.tmp_dir, "report.md"), encoding="utf-8") as fh:
            assert fh.read() == "# Old"  # server file untouched

    def test_atomic_failure_preserves_original(self, ctx, monkeypatch):
        _write_output(ctx, "report.md", "# Old")
        app = ctx.app

        def boom(path, content):
            raise OSError("simulated disk failure")

        monkeypatch.setattr(documents_module, "_atomic_write", boom)
        with app.test_client() as client:
            resp = client.put(
                f"/api/conversations/{ctx.conv_id}/documents/output:report.md",
                json={"content": "# Should not persist", "expected_hash": _hash_sha256("# Old")},
            )
        assert resp.status_code == 500
        with open(os.path.join(ctx.tmp_dir, "report.md"), encoding="utf-8") as fh:
            assert fh.read() == "# Old"

    def test_non_markdown_edit_rejected(self, ctx):
        _write_output(ctx, "notes.txt", "plain")
        app = ctx.app
        with app.test_client() as client:
            resp = client.put(
                f"/api/conversations/{ctx.conv_id}/documents/output:notes.txt",
                json={"content": "x", "expected_hash": _hash_sha256("plain")},
            )
        assert resp.status_code == 400

    def test_oversized_edit_rejected(self, ctx):
        _write_output(ctx, "report.md", "# Old")
        app = ctx.app
        big = "x" * (documents_module.MAX_EDITABLE_BYTES + 10)
        with app.test_client() as client:
            resp = client.put(
                f"/api/conversations/{ctx.conv_id}/documents/output:report.md",
                json={"content": big, "expected_hash": _hash_sha256("# Old")},
            )
        assert resp.status_code == 413
        with open(os.path.join(ctx.tmp_dir, "report.md"), encoding="utf-8") as fh:
            assert fh.read() == "# Old"

    def test_editing_disabled_when_local_access_off(self, tmp_db):
        app = _app(tmp_db)
        c = _make(app, local_access=False)
        _write_output(c, "report.md", "# Old")
        with app.test_client() as client:
            resp = client.get(f"/api/conversations/{c.conv_id}/documents")
        doc = next(d for d in resp.get_json()["documents"] if d["name"] == "report.md")
        assert doc["editable"] is False
        with app.test_client() as client:
            resp = client.put(
                f"/api/conversations/{c.conv_id}/documents/output:report.md",
                json={"content": "# Nope", "expected_hash": _hash_sha256("# Old")},
            )
        assert resp.status_code == 403


class TestUpdateUpload:
    def test_updates_existing_record_and_invalidates_cache(self, ctx):
        app = ctx.app
        up = _add_upload(ctx, "gen.md", "# V1")
        before = len(db_module.list_conv_files(ctx.conv_id))

        # Prime the extraction cache.
        from file_handler import _cached_upload_read
        with app.app_context():
            rec = db_module.get_conv_file(up["id"])
            _cached_upload_read(rec)
            assert file_handler._upload_content_cache.get(up["disk_path"]) is not None

        with app.test_client() as client:
            resp = client.put(
                f"/api/conversations/{ctx.conv_id}/documents/upload:{up['id']}",
                json={"content": "# V2 longer", "expected_hash": _hash_sha256("# V1")},
            )
        assert resp.status_code == 200

        with app.app_context():
            after = db_module.get_conv_file(up["id"])
            assert after["size_bytes"] == len("# V2 longer")
            assert after["char_count"] == len("# V2 longer")
            assert after["snippet"] == "# V2 longer"
            assert after["original_name"] == "gen.md"
            assert len(db_module.list_conv_files(ctx.conv_id)) == before  # no duplicate
            # Cache was evicted.
            assert file_handler._upload_content_cache.get(up["disk_path"]) is None
            # Disk content updated for the model via read_named_file.
            text, _ = file_handler.extract_text(after["disk_path"], "gen.md")
            assert text == "# V2 longer"


class TestManagedArtifacts:
    def _setup_project_conv(self, tmp_db):
        app = _app(tmp_db)
        workspace = tempfile.mkdtemp()
        _AUTO_CLEANUP.append(workspace)
        with app.app_context():
            db_module.set_setting("allow_local_file_access", "1")
            pid = db_module.create_folder("SDLC")["id"]
            db_module.register_project_from_template(pid, "sdlc", workspace)
            conv = db_module.create_conversation("BA", "m", folder_id=pid)
            db_module.update_conversation(conv["id"], output_dir=os.path.join(workspace, "Requirements"))
            conv_id = conv["id"]
        return app, pid, workspace, conv_id

    def test_artifact_marked_managed_in_list(self, tmp_db):
        app, pid, workspace, conv_id = self._setup_project_conv(tmp_db)
        with app.test_client() as client:
            resp = client.get(f"/api/conversations/{conv_id}/documents")
        docs = resp.get_json()["documents"]
        ba = next((d for d in docs if d["name"] == "BA-REQ.md"), None)
        assert ba is not None
        assert ba["managed_artifact"] is True
        assert ba["artifact_key"] == "BA-REQ"
        assert ba["project_id"] == pid

    def test_artifact_edit_routes_through_human_workflow(self, tmp_db):
        app, pid, workspace, conv_id = self._setup_project_conv(tmp_db)
        with app.app_context():
            artifact_before = db_module.get_artifact(pid, "BA-REQ")
            old = artifact_before["content_hash"]

        with open(os.path.join(workspace, "Requirements", "BA-REQ.md"), encoding="utf-8") as fh:
            old_content = fh.read()
        doc_id = "output:BA-REQ.md"
        new_content = old_content + "\n## REQ-014 — Edited requirement\n\nNew body.\n"

        with app.test_client() as client:
            resp = client.put(
                f"/api/conversations/{conv_id}/documents/{doc_id}",
                json={"content": new_content, "expected_hash": _hash_sha256(old_content)},
            )
        assert resp.status_code == 200
        assert resp.get_json()["managed_artifact"] is True

        with app.app_context():
            events = [e for e in db_module.list_change_events(pid) if e["source_key"] == "BA-REQ"]
            assert any(e["origin"] == "human" for e in events)
            artifact_after = db_module.get_artifact(pid, "BA-REQ")
            assert artifact_after["content_hash"] != old
            assert artifact_after["version"] == artifact_before["version"] + 1
            # Downstream artifacts were marked stale (review needed) but no
            # propagation jobs were auto-queued — only existing routes run those.
            all_jobs = []
            for event in db_module.list_change_events(pid):
                all_jobs.extend(db_module.list_propagation_jobs(event["id"]))
            assert all_jobs == []
        with open(os.path.join(workspace, "Requirements", "BA-REQ.md"), encoding="utf-8") as fh:
            assert fh.read() == new_content


class TestSseDocumentCreated:
    def _run_chat(self, tmp_db, config_output_dir, tool_name, path, content):
        from tests.fakes import FakeOpenAIClient

        app = _app(tmp_db)
        with app.app_context():
            db_module.set_setting("allow_local_file_access", "1")
            db_module.set_setting("output_dir", config_output_dir)
            endpoint = db_module.create_endpoint(
                "fake", "http://fase.openai.test/v1", api_key="f",
                default_model="m", is_default=True,
            )
            conv = db_module.create_conversation("C", "m", endpoint_id=endpoint["id"])
        fake = FakeOpenAIClient()
        fake.add_tool_calls_response([{
            "name": tool_name,
            "arguments": json.dumps({"path": path, "content": content}),
        }])
        fake.add_text_response("Done.")
        import app as app_module
        app_module.get_client = lambda endpoint: fake

        import io
        with app.test_client() as client:
            raw = client.post(
                f"/api/conversations/{conv['id']}/chat",
                json={"message": "write the file please"},
            ).get_data(as_text=True)
        events = []
        for line in raw.splitlines():
            line = line.strip()
            if not line.startswith("data: "):
                continue
            events.append(json.loads(line[6:]))
        return events

    def test_markdown_write_emits_document_created(self, tmp_db):
        out = tempfile.mkdtemp()
        _AUTO_CLEANUP.append(out)
        events = self._run_chat(tmp_db, out, "write_file", "report.md", "# Report")
        created = [e for e in events if e["type"] == "document_created"]
        assert created, [t for t in (e.get("type") for e in events)]
        assert created[0]["document"]["id"] == "output:report.md"
        assert created[0]["document"]["name"] == "report.md"
        assert created[0]["document"]["kind"] == "output"

    def test_non_markdown_write_no_event(self, tmp_db):
        out = tempfile.mkdtemp()
        _AUTO_CLEANUP.append(out)
        events = self._run_chat(tmp_db, out, "write_file", "data.txt", "plain")
        assert not [e for e in events if e["type"] == "document_created"]


class TestWriteFileRefreshMeta:
    def test_local_access_save_returns_document(self, tmp_db):
        app = _app(tmp_db)
        out = tempfile.mkdtemp()
        _AUTO_CLEANUP.append(out)
        with app.app_context():
            db_module.set_setting("allow_local_file_access", "1")
            db_module.set_setting("output_dir", out)
        with app.test_client() as client:
            resp = client.post("/api/write_file",
                               json={"path": "report.md", "content": "# X"})
        data = resp.get_json()
        assert data["success"] is True
        assert data["document"]["id"] == "output:report.md"
        assert data["document"]["kind"] == "output"

    def test_locked_save_returns_upload_document_and_creates_conv_file(self, tmp_db, monkeypatch):
        app = _app(tmp_db)
        upload_root = tempfile.mkdtemp()
        _AUTO_CLEANUP.append(upload_root)
        monkeypatch.setattr(file_handler, "UPLOAD_ROOT", upload_root)
        with app.app_context():
            db_module.set_setting("allow_local_file_access", "0")
            conv = db_module.create_conversation("C", "m")
        with app.test_client() as client:
            resp = client.post("/api/write_file", json={
                "path": "docs.md", "content": "# Locked mode",
                "conversation_id": conv["id"],
            })
        data = resp.get_json()
        assert data["success"] is True
        assert data["document"]["kind"] == "upload"
        assert data["document"]["id"].startswith("upload:")
        with app.app_context():
            files = db_module.list_conv_files(conv["id"])
        assert len(files) == 1
        assert files[0]["original_name"] == "docs.md"


class TestRenameAndBackups:
    def test_rename_output_document(self, ctx):
        _write_output(ctx, "report.md", "# Old")
        app = ctx.app
        with app.test_client() as client:
            resp = client.post(
                f"/api/conversations/{ctx.conv_id}/documents/output:report.md/rename",
                json={"name": "renamed.md"},
            )
        assert resp.status_code == 200
        data = resp.get_json()
        assert data["id"] == "output:renamed.md"
        assert not os.path.exists(os.path.join(ctx.tmp_dir, "report.md"))
        assert os.path.exists(os.path.join(ctx.tmp_dir, "renamed.md"))
        with app.test_client() as client:
            resp = client.get(f"/api/conversations/{ctx.conv_id}/documents/output:renamed.md")
        assert resp.status_code == 200
        assert resp.get_json()["content"] == "# Old"

    def test_rename_upload_updates_record_only(self, ctx):
        app = ctx.app
        up = _add_upload(ctx, "gen.md", "# Gen")
        disk_before = db_module.get_conv_file(up["id"])["disk_path"]
        with app.test_client() as client:
            resp = client.post(
                f"/api/conversations/{ctx.conv_id}/documents/upload:{up['id']}/rename",
                json={"name": "final.md"},
            )
        assert resp.status_code == 200
        assert resp.get_json()["id"] == f"upload:{up['id']}"
        with app.app_context():
            rec = db_module.get_conv_file(up["id"])
        assert rec["original_name"] == "final.md"
        assert rec["disk_path"] == disk_before  # uuid disk file untouched
        assert os.path.isfile(disk_before)
        assert len(db_module.list_conv_files(ctx.conv_id)) == 1

    def test_rename_rejects_unsafe_names(self, ctx):
        _write_output(ctx, "report.md", "# Old")
        app = ctx.app
        with app.test_client() as client:
            for bad in ("../escape.md", "sub/nested.md", "back\\slash.md", "noext", ".hidden.md"):
                resp = client.post(
                    f"/api/conversations/{ctx.conv_id}/documents/output:report.md/rename",
                    json={"name": bad},
                )
                assert resp.status_code in (400, 404), bad

    def test_rename_collision_409(self, ctx):
        _write_output(ctx, "report.md", "# A")
        _write_output(ctx, "other.md", "# B")
        app = ctx.app
        with app.test_client() as client:
            resp = client.post(
                f"/api/conversations/{ctx.conv_id}/documents/output:report.md/rename",
                json={"name": "other.md"},
            )
        assert resp.status_code == 409

    def test_rename_managed_artifact_rejected(self, tmp_db):
        app = _app(tmp_db)
        workspace = tempfile.mkdtemp()
        _AUTO_CLEANUP.append(workspace)
        with app.app_context():
            db_module.set_setting("allow_local_file_access", "1")
            pid = db_module.create_folder("SDLC")["id"]
            db_module.register_project_from_template(pid, "sdlc", workspace)
            conv = db_module.create_conversation("BA", "m", folder_id=pid)
            db_module.update_conversation(conv["id"], output_dir=os.path.join(workspace, "Requirements"))
        with app.test_client() as client:
            resp = client.post(
                f"/api/conversations/{conv['id']}/documents/output:BA-REQ.md/rename",
                json={"name": "moved.md"},
            )
        assert resp.status_code == 409

    def test_output_edits_keep_limited_backups(self, ctx):
        _write_output(ctx, "report.md", "# v1")
        app = ctx.app
        with app.test_client() as client:
            client.put(
                f"/api/conversations/{ctx.conv_id}/documents/output:report.md",
                json={"content": "# v2", "expected_hash": _hash_sha256("# v1")},
            )
        backup_dir = os.path.join(ctx.tmp_dir, ".agents", "editor-backups")
        assert os.path.isdir(backup_dir)
        backups = [f for f in os.listdir(backup_dir) if f.endswith(".md")]
        assert len(backups) == 1
        with open(os.path.join(backup_dir, backups[0]), encoding="utf-8") as fh:
            assert fh.read() == "# v1"

    def test_upload_edits_do_not_create_backups(self, ctx):
        app = ctx.app
        up = _add_upload(ctx, "gen.md", "# v1")
        with app.test_client() as client:
            client.put(
                f"/api/conversations/{ctx.conv_id}/documents/upload:{up['id']}",
                json={"content": "# v2", "expected_hash": _hash_sha256("# v1")},
            )
        assert not os.path.exists(os.path.join(ctx.tmp_dir, ".agents"))


class TestImportFromOutput:
    def test_browse_starts_at_output_folder_root(self, ctx):
        _write_output(ctx, "report.md", "# Report")
        _write_output(ctx, "sub/notes.markdown", "# Notes")
        _write_output(ctx, "sub/readme.txt", "plain")
        _write_output(ctx, "code.py", "print(1)")
        app = ctx.app
        with app.test_client() as client:
            resp = client.get(f"/api/conversations/{ctx.conv_id}/import-browse")
        assert resp.status_code == 200
        data = resp.get_json()
        assert data["current"] == ""
        assert data["up"] is None
        names = {e["name"]: e for e in data["entries"]}
        assert "report.md" in names and not names["report.md"]["is_dir"]
        assert names["report.md"]["ext"] == ".md"
        assert "sub" in names and names["sub"]["is_dir"]
        assert "code.py" not in names  # non-markdown/txt hidden

    def test_browse_navigates_subfolder_and_up(self, ctx):
        _write_output(ctx, "sub/inner.md", "# Inner")
        app = ctx.app
        with app.test_client() as client:
            resp = client.get(f"/api/conversations/{ctx.conv_id}/import-browse?path=sub")
            assert resp.status_code == 200
            data = resp.get_json()
            assert data["current"] == "sub"
            assert data["up"] == ""
            assert [e["name"] for e in data["entries"]] == ["inner.md"]

    def test_browse_rejects_traversal(self, ctx):
        app = ctx.app
        with app.test_client() as client:
            resp = client.get(f"/api/conversations/{ctx.conv_id}/import-browse?path=..%2F..%2F..")
        assert resp.status_code == 400

    def test_import_creates_conv_file_and_document(self, ctx):
        _write_output(ctx, "report.md", "# Report\n\nhello")
        app = ctx.app
        with app.test_client() as client:
            resp = client.post(
                f"/api/conversations/{ctx.conv_id}/import-from-output",
                json={"path": "report.md"},
            )
        assert resp.status_code == 201
        data = resp.get_json()
        assert data["original_name"] == "report.md"
        assert data["success"] is True
        assert data["document"] and data["document"]["kind"] == "upload"
        with app.app_context():
            files = db_module.list_conv_files(ctx.conv_id)
        assert len(files) == 1
        with open(files[0]["disk_path"], encoding="utf-8") as fh:
            assert fh.read() == "# Report\n\nhello"
        # Original output file remains untouched (copy, not move).
        with open(os.path.join(ctx.tmp_dir, "report.md"), encoding="utf-8") as fh:
            assert fh.read() == "# Report\n\nhello"
        # Available to the model through read_named_file next turn.
        with app.app_context():
            text, _ = file_handler.extract_text(files[0]["disk_path"], "report.md")
        assert "hello" in text
        # The upload is exposed via the conversation files API too.
        with app.test_client() as client:
            files_api = client.get(f"/api/conversations/{ctx.conv_id}/files").get_json()
        assert any(f["original_name"] == "report.md" for f in files_api)

    def test_import_txt_allowed(self, ctx):
        _write_output(ctx, "data/raw.txt", "plain text")
        app = ctx.app
        with app.test_client() as client:
            resp = client.post(
                f"/api/conversations/{ctx.conv_id}/import-from-output",
                json={"path": "data/raw.txt"},
            )
        assert resp.status_code == 201
        assert resp.get_json()["original_name"] == "raw.txt"

    def test_import_rejects_missing_and_non_text(self, ctx):
        _write_output(ctx, "code.py", "print(1)")
        app = ctx.app
        with app.test_client() as client:
            resp = client.post(
                f"/api/conversations/{ctx.conv_id}/import-from-output",
                json={"path": "missing.md"},
            )
            assert resp.status_code == 404
            resp = client.post(
                f"/api/conversations/{ctx.conv_id}/import-from-output",
                json={"path": "code.py"},
            )
            assert resp.status_code == 415

    def test_import_without_output_dir_400(self, tmp_db):
        app = _app(tmp_db)
        with app.app_context():
            db_module.set_setting("output_dir", "")
            conv = db_module.create_conversation("C", "m")
        with app.test_client() as client:
            resp = client.get(f"/api/conversations/{conv['id']}/import-browse")
        assert resp.status_code == 400


class TestRecentEditNote:
    def test_note_after_save(self, ctx):
        _write_output(ctx, "report.md", "# Old")
        app = ctx.app
        with app.test_client() as client:
            client.put(
                f"/api/conversations/{ctx.conv_id}/documents/output:report.md",
                json={"content": "# New", "expected_hash": _hash_sha256("# Old")},
            )
        with app.app_context():
            note = db_module.get_recent_document_edit_note(ctx.conv_id)
        assert "report.md" in note
        assert "read_named_file" not in note  # output kind → read_file hint
        assert "read_file" in note

        with app.app_context():
            conv = db_module.get_conversation(ctx.conv_id)
            from app import _build_system_prompt
            system, _ = _build_system_prompt(conv, ctx.conv_id, tools_on=True)
        assert "edited report.md" in system

    def test_upload_note_hints_read_named_file(self, ctx):
        app = ctx.app
        up = _add_upload(ctx, "notes.md", "# V1")
        with app.test_client() as client:
            client.put(
                f"/api/conversations/{ctx.conv_id}/documents/upload:{up['id']}",
                json={"content": "# V2", "expected_hash": _hash_sha256("# V1")},
            )
        with app.app_context():
            note = db_module.get_recent_document_edit_note(ctx.conv_id)
        assert "read_named_file" in note


class TestDocxExport:
    DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

    def _reopen(self, payload):
        from docx import Document as DocxDocument
        from io import BytesIO
        return DocxDocument(BytesIO(payload))

    def test_export_output_document_docx(self, ctx):
        _write_output(ctx, "report.md", "# Hello\n\nThis is **bold** text.\n")
        app = ctx.app
        with app.test_client() as client:
            resp = client.get(
                f"/api/conversations/{ctx.conv_id}/documents/output:report.md/export?format=docx"
            )
        assert resp.status_code == 200
        assert resp.mimetype == self.DOCX_MIME
        assert "filename=\"report.docx\"" in resp.headers.get("Content-Disposition", "")
        payload = resp.data
        assert payload[:2] == b"PK"  # ZIP/DOCX magic
        doc = self._reopen(payload)
        texts = [p.text for p in doc.paragraphs]
        assert any(t == "Hello" for t in texts)
        assert any("bold" in t for t in texts)

    def test_export_posts_local_content(self, ctx):
        _write_output(ctx, "report.md", "# Server version\n")
        app = ctx.app
        with app.test_client() as client:
            resp = client.post(
                f"/api/conversations/{ctx.conv_id}/documents/output:report.md/export?format=docx",
                json={"content": "## Edited locally\n", "name": "report.md"},
            )
        assert resp.status_code == 200
        doc = self._reopen(resp.data)
        assert any(p.text.strip() == "Edited locally" for p in doc.paragraphs)

    def test_export_rejects_unknown_format(self, ctx):
        _write_output(ctx, "report.md", "# Hi")
        app = ctx.app
        with app.test_client() as client:
            resp = client.get(
                f"/api/conversations/{ctx.conv_id}/documents/output:report.md/export?format=pdf"
            )
        assert resp.status_code == 400

    def test_export_bad_document_id(self, ctx):
        app = ctx.app
        with app.test_client() as client:
            resp = client.get(
                f"/api/conversations/{ctx.conv_id}/documents/output:../../etc/export?format=docx"
            )
        assert resp.status_code in (400, 404)


if __name__ == "__main__":
    pytest.main([__file__, "-v"])